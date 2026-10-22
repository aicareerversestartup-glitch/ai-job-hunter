"""Render ``CVData`` and cover letters to DOCX, PDF and Markdown (ATS-friendly single column)."""
from __future__ import annotations

import io

from backend.cv.schema import CVData


def _dates(start: str, end: str) -> str:
    return " – ".join(p for p in (start, end) if p)


def cv_to_markdown(cv: CVData) -> str:
    c = cv.contact
    out = [f"# {c.full_name}".rstrip()]
    contact_bits = [b for b in (c.email, c.phone, c.location, *c.links) if b]
    if contact_bits:
        out.append(" · ".join(contact_bits))
    if cv.headline:
        out.append(f"\n**{cv.headline}**")
    if cv.summary:
        out += ["\n## Summary", cv.summary]
    if cv.experience:
        out.append("\n## Experience")
        for e in cv.experience:
            out.append(f"\n### {e.title} — {e.employer}")
            meta = " · ".join(p for p in (_dates(e.start, e.end), e.location) if p)
            if meta:
                out.append(f"*{meta}*")
            out += [f"- {b}" for b in e.bullets]
    if cv.skills:
        out += ["\n## Skills", ", ".join(cv.skills)]
    if cv.projects:
        out.append("\n## Projects")
        for p in cv.projects:
            out.append(f"\n### {p.name}")
            if p.description:
                out.append(p.description)
            out += [f"- {b}" for b in p.bullets]
            if p.technologies:
                out.append(f"*Technologies:* {', '.join(p.technologies)}")
    if cv.education:
        out.append("\n## Education")
        for ed in cv.education:
            degree = " in ".join(p for p in (ed.degree, ed.field) if p)
            line = f"- **{degree or ed.institution}**" + (f", {ed.institution}" if degree else "")
            if _dates(ed.start, ed.end):
                line += f" ({_dates(ed.start, ed.end)})"
            out.append(line)
    if cv.certifications:
        out += ["\n## Certifications", *[f"- {x}" for x in cv.certifications]]
    if cv.languages:
        out += ["\n## Languages", ", ".join(cv.languages)]
    for title, lines in cv.other_sections.items():
        out += [f"\n## {title}", *[f"- {x}" for x in lines]]
    return "\n".join(out).strip() + "\n"


def cv_to_docx(cv: CVData) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(46)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(cv.contact.full_name)
    run.bold, run.font.size = True, Pt(18)
    bits = [b for b in (cv.contact.email, cv.contact.phone, cv.contact.location, *cv.contact.links) if b]
    if bits:
        p = doc.add_paragraph(" | ".join(bits))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cv.headline:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(cv.headline).italic = True

    def heading(text: str):
        h = doc.add_paragraph()
        r = h.add_run(text.upper())
        r.bold, r.font.size = True, Pt(11.5)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)

    def bullet(text: str):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    if cv.summary:
        heading("Summary")
        doc.add_paragraph(cv.summary)
    if cv.experience:
        heading("Experience")
        for e in cv.experience:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.add_run(f"{e.title}").bold = True
            p.add_run(f" — {e.employer}")
            meta = " · ".join(x for x in (_dates(e.start, e.end), e.location) if x)
            if meta:
                p.add_run(f"\n{meta}").italic = True
            for b in e.bullets:
                bullet(b)
    if cv.skills:
        heading("Skills")
        doc.add_paragraph(", ".join(cv.skills))
    if cv.projects:
        heading("Projects")
        for pr in cv.projects:
            p = doc.add_paragraph()
            p.add_run(pr.name).bold = True
            if pr.description:
                p.add_run(f" — {pr.description}")
            for b in pr.bullets:
                bullet(b)
            if pr.technologies:
                doc.add_paragraph("Technologies: " + ", ".join(pr.technologies))
    if cv.education:
        heading("Education")
        for ed in cv.education:
            degree = " in ".join(x for x in (ed.degree, ed.field) if x)
            p = doc.add_paragraph()
            p.add_run(degree or ed.institution).bold = True
            if degree:
                p.add_run(f", {ed.institution}")
            if _dates(ed.start, ed.end):
                p.add_run(f" ({_dates(ed.start, ed.end)})")
            for d in ed.details:
                bullet(d)
    if cv.certifications:
        heading("Certifications")
        for c in cv.certifications:
            bullet(c)
    if cv.languages:
        heading("Languages")
        doc.add_paragraph(", ".join(cv.languages))
    for title, lines in cv.other_sections.items():
        heading(title)
        for line in lines:
            bullet(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def text_to_docx(text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for block in text.split("\n\n"):
        doc.add_paragraph(block.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def cv_to_pdf(cv: CVData) -> bytes:
    return _markdownish_pdf(cv_to_markdown(cv))


def text_to_pdf(text: str) -> bytes:
    return _markdownish_pdf(text, plain=True)


def _markdownish_pdf(source: str, plain: bool = False) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=13.5)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=TA_CENTER)
    center = ParagraphStyle("center", parent=body, alignment=TA_CENTER, textColor="#444444")
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=11.5, spaceBefore=9, spaceAfter=3,
                        textColor="#1f3a5f")
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=5)

    def inline(text: str) -> str:
        import re
        text = escape(text)
        text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
        return re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)

    story = []
    pending: list = []

    def flush_bullets():
        nonlocal pending
        if pending:
            story.append(ListFlowable([ListItem(Paragraph(inline(b), body), leftIndent=10) for b in pending],
                                      bulletType="bullet", start="•", leftIndent=12))
            pending = []

    lines = source.splitlines()
    for index, line in enumerate(lines):
        stripped = line.strip()
        if plain:
            story.append(Paragraph(inline(stripped), body) if stripped else Spacer(1, 6))
            continue
        if stripped.startswith("- "):
            pending.append(stripped[2:])
            continue
        flush_bullets()
        if not stripped:
            continue
        if stripped.startswith("# "):
            story.append(Paragraph(inline(stripped[2:]), h1))
        elif stripped.startswith("## "):
            story.append(Paragraph(inline(stripped[3:]), h2))
        elif stripped.startswith("### "):
            story.append(Paragraph(inline(stripped[4:]), h3))
        elif index == 1:
            story.append(Paragraph(inline(stripped), center))
        else:
            story.append(Paragraph(inline(stripped), body))
    flush_bullets()
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm,
                      bottomMargin=15 * mm).build(story)
    return buffer.getvalue()
