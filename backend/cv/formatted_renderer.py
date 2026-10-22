"""Render a tailored CV in a specific country format and visual style (DOCX and PDF).

Fitting to the country's page limit only REMOVES content (least relevant bullets of older roles
first, then projects). It never rewrites or adds, so the truthfulness guarantee still holds.
"""
from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from backend.cv.formats import CountryFormat
from backend.cv.profile import CandidateProfile
from backend.cv.schema import CVData
from backend.utils.dates import parse_cv_date, total_experience_months

DEFAULT_HEADINGS = {
    "experience": "Professional Experience", "education": "Education", "skills": "Key Skills",
    "certifications": "Certifications", "languages": "Languages", "projects": "Projects", "referees": "Referees",
    "other": "Additional Information",
}
LINES_PER_PAGE = 54
CHARS_PER_LINE = 100


@dataclass
class RenderedCV:
    name: str
    title_line: str
    contact_line: str
    personal_lines: list[str]
    photo: bytes | None
    sections: list[tuple[str, str, object]]  # (key, heading, payload)
    signature: str = ""
    removed: list[str] = field(default_factory=list)  # what page-fitting left out (shown to the user)


def _fmt_date(value: str, style: str, precision: str = "month") -> str:
    if value and precision == "year" and value.strip().lower() not in {"present", "current", "now", "today"}:
        parsed = parse_cv_date(value)
        return str(parsed.year) if parsed else value
    if not value or style != "MM/YYYY":
        return value
    parsed = parse_cv_date(value)
    if parsed is None or value.strip().lower() in {"present", "current", "now", "today"}:
        return value
    if value.strip().isdigit():  # year only stays a year
        return value
    return parsed.strftime("%m/%Y")


def effective_max_pages(cv: CVData, fmt: CountryFormat) -> int:
    if fmt.code == "US":
        years = total_experience_months([(e.start, e.end) for e in cv.experience]) / 12
        return 1 if years < 10 else 2
    return fmt.max_pages


def _estimate_lines(cv: CVData, fmt: CountryFormat, scale: float) -> int:
    def wrap(text: str) -> int:
        return max(1, math.ceil(len(text) * scale / CHARS_PER_LINE))

    lines = 6
    sections = set(fmt.section_order)
    if "summary" in sections and cv.summary:
        lines += 2 + wrap(cv.summary)
    if "skills" in sections and cv.skills:
        lines += 2 + wrap(", ".join(cv.skills))
    if "experience" in sections:
        lines += 2
        for e in cv.experience:
            lines += 2 + sum(wrap(b) for b in e.bullets)
    if "projects" in sections and cv.projects:
        lines += 2 + sum(2 + sum(wrap(b) for b in p.bullets) for p in cv.projects)
    if "education" in sections:
        lines += 2 + 2 * len(cv.education)
    if "certifications" in sections and cv.certifications:
        lines += 2 + len(cv.certifications)
    if "languages" in sections and cv.languages:
        lines += 3
    return lines


def fit_to_pages(cv: CVData, fmt: CountryFormat, style: dict, factor: float = 1.0) -> tuple[CVData, list[str]]:
    fitted = cv.model_copy(deep=True)
    removed: list[str] = []
    scale = style.get("size", 10.5) / 10.5 * style.get("spacing", 1.0)
    budget = effective_max_pages(cv, fmt) * LINES_PER_PAGE * factor / scale
    while _estimate_lines(fitted, fmt, scale) > budget:
        # Oldest role first; keep at least two bullets per role (bullets are already relevance-sorted).
        candidates = [e for e in reversed(fitted.experience) if len(e.bullets) > 2]
        if candidates:
            role = candidates[0]
            removed.append(f"{role.title}, {role.employer}: {role.bullets.pop()}")
            continue
        if fitted.projects:
            removed.append(f"Project: {fitted.projects.pop().name}")
            continue
        break
    return fitted, removed


def _pull_achievements(cv: CVData, limit: int = 5) -> list[str]:
    """Hybrid layout: move up to five quantified bullets (already relevance-sorted) to the top, each
    labelled with its employer, spreading picks across employers. Every role keeps at least one bullet."""
    picked: list[str] = []
    progress = True
    while len(picked) < limit and progress:
        progress = False
        for exp in cv.experience:
            if len(picked) >= limit:
                break
            candidate = next((b for b in exp.bullets if any(ch.isdigit() for ch in b)), None)
            if candidate and len(exp.bullets) > 1:
                exp.bullets.remove(candidate)
                picked.append(f"{candidate.rstrip('.')} ({exp.employer}).")
                progress = True
    return picked


def _group_roles(cv: CVData, dates) -> list[dict]:
    """Consecutive roles at the same employer are shown as promotions under one employer."""
    from backend.jobs.normalizer import normalize_company

    groups: list[dict] = []
    for e in cv.experience:
        label = {"contract": "Contract", "consulting": "Consulting Assignment", "freelance": "Freelance",
                 "internship": "Internship"}.get((e.employment_type or "").lower(), "")
        title = f"{e.title} ({label})" if label else e.title
        role = {"title": title, "dates": dates(e.start, e.end), "bullets": e.bullets, "location": e.location}
        if groups and normalize_company(groups[-1]["employer"]) == normalize_company(e.employer):
            g = groups[-1]
            g["roles"].append(role)
            g["start"] = e.start
            g["dates"] = dates(g["start"], g["end"])
        else:
            groups.append({"employer": e.employer, "location": e.location, "start": e.start, "end": e.end,
                           "dates": dates(e.start, e.end), "roles": [role]})
    return groups


def mobility_lines(profile: CandidateProfile, job_country: str) -> list[str]:
    """'Global mobility' lines, only for jobs in countries where you can't already work."""
    from backend.utils.text import normalize_key

    if not job_country:
        return []
    authorized = {normalize_key(c) for c in profile.work_authorized_countries}
    if normalize_key(job_country) in authorized:
        return []
    lines = []
    if profile.relocation_statement:
        lines.append(profile.relocation_statement)
    if profile.passport_status:
        lines.append(f"Passport: {profile.passport_status}")
    return lines


def relocation_line(profile: CandidateProfile, job_country: str) -> str:
    lines = mobility_lines(profile, job_country)
    return lines[0] if lines and profile.relocation_statement else ""


def build(cv: CVData, fmt: CountryFormat, style: dict, profile: CandidateProfile, factor: float = 1.0,
          job_country: str = "") -> RenderedCV:
    fitted, removed = fit_to_pages(cv, fmt, style, factor)
    headings = {**DEFAULT_HEADINGS, "summary": fmt.summary_heading, **style.get("headings", {})}
    c = fitted.contact
    contact = [x for x in (c.email, c.phone, c.location, *c.links) if x]
    personal: list[str] = []
    details = {"dob": ("Date of birth", profile.date_of_birth), "nationality": ("Nationality", profile.nationality),
               "visa_status": ("Visa status", profile.visa_status),
               "driving_licence": ("Driving licence", profile.driving_licence)}
    for key in fmt.personal_details:
        label, value = details.get(key, ("", ""))
        if value:
            personal.append(f"{label}: {value}")
    personal += mobility_lines(profile, job_country)
    photo = None
    if fmt.photo in ("expected", "common") and profile.include_photo and profile.photo_path:
        path = Path(profile.photo_path)
        if path.exists():
            photo = path.read_bytes()

    precision = style.get("date_precision", "month")

    def dates(start: str, end: str) -> str:
        return " – ".join(x for x in (_fmt_date(start, fmt.date_style, precision),
                                      _fmt_date(end, fmt.date_style, precision)) if x)

    hybrid = style.get("layout") == "hybrid"
    achievements = _pull_achievements(fitted) if hybrid else []
    order = list(fmt.section_order)
    if hybrid and "certifications" in order and fitted.certifications:
        # Credentials right after the summary/achievements block.
        order.remove("certifications")
        anchor = order.index("summary") + 1 if "summary" in order else 0
        order.insert(anchor, "certifications")
    sections: list[tuple[str, str, object]] = []
    if achievements and "summary" not in fmt.section_order:
        sections.append(("achievements", "Key Achievements", achievements))
    for key in order:
        if key == "summary" and fitted.summary:
            sections.append((key, headings[key], fitted.summary))
            if achievements:
                sections.append(("achievements", "Key Achievements", achievements))
        elif key == "summary" and achievements:
            sections.append(("achievements", "Key Achievements", achievements))
        elif key == "skills" and fitted.skills:
            sections.append((key, headings[key], fitted.skills))
        elif key == "experience" and fitted.experience:
            sections.append((key, headings[key], _group_roles(fitted, dates)))
        elif key == "projects" and fitted.projects:
            sections.append((key, headings[key], [{"name": p.name, "description": p.description, "bullets": p.bullets,
                                                   "technologies": p.technologies} for p in fitted.projects]))
        elif key == "education" and fitted.education:
            items = [{"degree": " in ".join(x for x in (e.degree, e.field) if x), "institution": e.institution,
                      "dates": dates(e.start, e.end), "details": e.details} for e in fitted.education]
            sections.append((key, headings[key], items))
        elif key == "certifications" and fitted.certifications:
            sections.append((key, headings[key], fitted.certifications))
        elif key == "languages" and fitted.languages:
            sections.append((key, headings[key], fitted.languages))
        elif key == "other" and fitted.other_sections:
            for title, lines in fitted.other_sections.items():
                sections.append(("other", title, lines))
        elif key == "referees" and fmt.referees == "list" and profile.referees:
            sections.append((key, headings[key], profile.referees))
    signature = ""
    if fmt.signature_line:
        place = (profile.location or c.location).split(",")[0]
        signature = f"{place + ', ' if place else ''}{date.today().strftime('%d.%m.%Y')}\n" + (
            c.full_name.title() if c.full_name.isupper() else c.full_name)
    display_name = c.full_name.title() if c.full_name.isupper() else c.full_name
    return RenderedCV(name=display_name, title_line=fitted.headline, contact_line=" | ".join(contact),
                      personal_lines=personal, photo=photo, sections=sections, signature=signature, removed=removed)


# --------------------------------------------------------------------------------------------- DOCX
def to_docx(doc: RenderedCV, style: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    accent = RGBColor.from_string(style["accent"].lstrip("#"))
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = style["font"]
    normal.font.size = Pt(style["size"])
    normal.paragraph_format.space_after = Pt(2 * style["spacing"])
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.6)
    align = WD_ALIGN_PARAGRAPH.CENTER if style["align_name"] == "center" else WD_ALIGN_PARAGRAPH.LEFT

    if doc.photo:
        table = d.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        target_par = left.paragraphs[0]
        try:
            right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right.paragraphs[0].add_run().add_picture(io.BytesIO(doc.photo), width=Inches(1.1))
        except Exception:
            pass
        header_container = left
    else:
        target_par = d.add_paragraph()
        header_container = None
    target_par.alignment = align
    r = target_par.add_run(doc.name)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(style["size"] + 8), accent

    def header_line(text: str, italic: bool = False):
        p = header_container.add_paragraph() if header_container is not None else d.add_paragraph()
        p.alignment = align
        run = p.add_run(text)
        run.italic = italic
        run.font.size = Pt(style["size"] - 0.5)

    if doc.title_line:
        header_line(doc.title_line, italic=True)
    if doc.contact_line:
        header_line(doc.contact_line)
    for line in doc.personal_lines:
        header_line(line)

    def heading(text: str):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(9 * style["spacing"])
        run = p.add_run(text)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(style["size"] + 1.5), accent
        _bottom_border(p, style["accent"].lstrip("#"))

    def bullet(text: str):
        p = d.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    for key, title, payload in doc.sections:
        heading(title)
        if key == "summary":
            d.add_paragraph(str(payload))
        elif key == "skills":
            d.add_paragraph("  •  ".join(payload))  # type: ignore[arg-type]
        elif key == "experience":
            for group in payload:  # type: ignore[union-attr]
                multi = len(group["roles"]) > 1
                if multi:
                    p = d.add_paragraph()
                    p.paragraph_format.space_before = Pt(5)
                    p.add_run(group["employer"]).bold = True
                    meta = "  |  ".join(x for x in (group["dates"], group["location"]) if x)
                    if meta:
                        p.add_run(f"  |  {meta}").font.size = Pt(style["size"] - 1)
                for item in group["roles"]:
                    p = d.add_paragraph()
                    p.paragraph_format.space_before = Pt(3 if multi else 5)
                    r = p.add_run(item["title"])
                    r.bold, r.italic = not multi, multi
                    if not multi:
                        p.add_run(f"  |  {group['employer']}")
                    meta = "  |  ".join(x for x in (item["dates"], "" if multi else item["location"]) if x)
                    if meta:
                        m = p.add_run(f"\n{meta}" if not multi else f"  ({meta})")
                        m.italic, m.font.size = True, Pt(style["size"] - 1)
                    for b in item["bullets"]:
                        bullet(b)
        elif key == "projects":
            for item in payload:  # type: ignore[union-attr]
                p = d.add_paragraph()
                p.add_run(item["name"]).bold = True
                if item["description"]:
                    p.add_run(f" — {item['description']}")
                for b in item["bullets"]:
                    bullet(b)
        elif key == "education":
            for item in payload:  # type: ignore[union-attr]
                p = d.add_paragraph()
                p.add_run(item["degree"] or item["institution"]).bold = True
                if item["degree"]:
                    p.add_run(f", {item['institution']}")
                if item["dates"]:
                    p.add_run(f"  ({item['dates']})")
                for x in item["details"]:
                    bullet(x)
        elif key == "languages":
            d.add_paragraph(", ".join(payload))  # type: ignore[arg-type]
        elif key == "referees":
            for ref in payload:  # type: ignore[union-attr]
                d.add_paragraph(", ".join(str(ref.get(k, "")) for k in ("name", "title", "company", "email", "phone",
                                                                        "relationship") if ref.get(k)))
        else:
            for line in payload:  # type: ignore[union-attr]
                bullet(str(line))
    if doc.signature:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.add_run(doc.signature)
    buffer = io.BytesIO()
    d.save(buffer)
    return buffer.getvalue()


def _bottom_border(paragraph, color_hex: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color_hex)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


# ---------------------------------------------------------------------------------------------- PDF
def to_pdf(doc: RenderedCV, style: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (HRFlowable, Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    size, sp = style["size"], style["spacing"]
    accent = HexColor(style["accent"])
    align = TA_CENTER if style["align_name"] == "center" else TA_LEFT
    body = ParagraphStyle("b", fontName=style["pdf_font"], fontSize=size, leading=size * 1.35 * sp)
    small = ParagraphStyle("s", parent=body, fontSize=size - 1, textColor=HexColor("#555555"), alignment=align)
    name_style = ParagraphStyle("n", parent=body, fontName=style["pdf_bold"], fontSize=size + 8,
                                leading=(size + 8) * 1.2, textColor=accent, alignment=align)
    h = ParagraphStyle("h", parent=body, fontName=style["pdf_bold"], fontSize=size + 1.5, textColor=accent,
                       spaceBefore=8 * sp, spaceAfter=1)
    role = ParagraphStyle("r", parent=body, spaceBefore=4 * sp)
    e = lambda t: escape(str(t))  # noqa: E731

    header = [Paragraph(e(doc.name), name_style)]
    if doc.title_line:
        header.append(Paragraph(f"<i>{e(doc.title_line)}</i>", small))
    if doc.contact_line:
        header.append(Paragraph(e(doc.contact_line), small))
    for line in doc.personal_lines:
        header.append(Paragraph(e(line), small))
    story: list = []
    if doc.photo:
        try:
            img = Image(io.BytesIO(doc.photo), width=28 * mm, height=35 * mm)
            t = Table([[header, img]], colWidths=[140 * mm, 32 * mm])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
            story.append(t)
        except Exception:
            story += header
    else:
        story += header

    def bullets(items):
        return ListFlowable([ListItem(Paragraph(e(b), body), leftIndent=10) for b in items], bulletType="bullet",
                            start="•", leftIndent=12)

    for key, title, payload in doc.sections:
        story += [Paragraph(e(title), h), HRFlowable(width="100%", thickness=0.6, color=accent, spaceAfter=3)]
        if key == "summary":
            story.append(Paragraph(e(payload), body))
        elif key == "skills":
            story.append(Paragraph(" &nbsp;•&nbsp; ".join(e(s) for s in payload), body))  # type: ignore[union-attr]
        elif key == "experience":
            for group in payload:  # type: ignore[union-attr]
                multi = len(group["roles"]) > 1
                if multi:
                    meta = " | ".join(x for x in (group["dates"], group["location"]) if x)
                    story.append(Paragraph(f"<b>{e(group['employer'])}</b>"
                                           + (f" <font size='{size - 1}' color='#555555'>| {e(meta)}</font>" if meta else ""), role))
                for item in group["roles"]:
                    if multi:
                        story.append(Paragraph(f"<i>{e(item['title'])}</i> <font size='{size - 1}' color='#555555'>"
                                               f"({e(item['dates'])})</font>", body))
                    else:
                        meta = " | ".join(x for x in (item["dates"], item["location"]) if x)
                        story.append(Paragraph(f"<b>{e(item['title'])}</b> | {e(group['employer'])}"
                                               + (f"<br/><font size='{size - 1}' color='#555555'><i>{e(meta)}</i></font>" if meta else ""),
                                               role))
                    if item["bullets"]:
                        story.append(bullets(item["bullets"]))
        elif key == "projects":
            for item in payload:  # type: ignore[union-attr]
                story.append(Paragraph(f"<b>{e(item['name'])}</b>" + (f" — {e(item['description'])}" if item["description"] else ""), role))
                if item["bullets"]:
                    story.append(bullets(item["bullets"]))
        elif key == "education":
            for item in payload:  # type: ignore[union-attr]
                text = f"<b>{e(item['degree'] or item['institution'])}</b>"
                if item["degree"]:
                    text += f", {e(item['institution'])}"
                if item["dates"]:
                    text += f" ({e(item['dates'])})"
                story.append(Paragraph(text, role))
                if item["details"]:
                    story.append(bullets(item["details"]))
        elif key == "languages":
            story.append(Paragraph(e(", ".join(payload)), body))  # type: ignore[arg-type]
        elif key == "referees":
            for ref in payload:  # type: ignore[union-attr]
                story.append(Paragraph(e(", ".join(str(ref.get(k, "")) for k in ("name", "title", "company", "email",
                                                                                "phone", "relationship") if ref.get(k))), body))
        else:
            story.append(bullets([str(x) for x in payload]))  # type: ignore[union-attr]
    if doc.signature:
        story += [Spacer(1, 18), Paragraph(e(doc.signature).replace("\n", "<br/>"), body)]
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=19 * mm, rightMargin=19 * mm, topMargin=15 * mm,
                      bottomMargin=15 * mm, title=f"{doc.name} CV").build(story)
    return buffer.getvalue()


def pdf_page_count(data: bytes) -> int:
    from pypdf import PdfReader

    return len(PdfReader(io.BytesIO(data)).pages)


def render(cv: CVData, fmt: CountryFormat, style: dict, profile: CandidateProfile,
           job_country: str = "") -> tuple[bytes, bytes, RenderedCV]:
    """Build DOCX + PDF, re-trimming until the PDF really fits the country's page limit."""
    limit = effective_max_pages(cv, fmt)
    factor = 1.0
    for _ in range(6):
        doc = build(cv, fmt, style, profile, factor, job_country)
        pdf = to_pdf(doc, style)
        if pdf_page_count(pdf) <= limit:
            break
        factor *= 0.88
    return to_docx(doc, style), pdf, doc
