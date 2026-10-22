"""Company search, JD breakdown, country CV formats, sample-layout learning, interview email tracking."""
import io
from datetime import datetime, timezone

from backend.cv.formats import FORMATS, STYLES, code_for_country, get_format, resolve_layout
from backend.cv.formatted_renderer import build, pdf_page_count, render
from backend.cv.layout_import import learn_layout
from backend.cv.truthfulness import check_cv
from backend.integrations.companies import company_matches, company_names, load_registry, slug_candidates
from backend.integrations.inbox import InboxMessage, classify, detect, match_application
from backend.jobs.jd_sections import split_jd
from tests.helpers import sample_cv, sample_facts, sample_job, sample_profile


# ------------------------------------------------------------------ companies
def test_registry_has_users_companies():
    names = {c["name"] for c in load_registry()}
    assert {"DHL Group", "Amazon", "Maersk", "Nestlé", "Schneider Electric"} <= names


def test_company_matching_uses_aliases_and_whole_words():
    names = company_names("DHL Group")
    assert company_matches("DHL Supply Chain (UK) Limited", names)
    assert company_matches("Deutsche Post DHL", names)
    assert not company_matches("Amazonia Foods", company_names("Amazon"))
    assert company_matches("Amazon Logistics", company_names("Amazon"))


def test_slug_candidates():
    assert "kuehnenagel" in slug_candidates("Kuehne+Nagel")
    assert "schneider-electric" in slug_candidates("Schneider Electric")


# ------------------------------------------------------------------ JD breakdown
def test_jd_sections():
    parts = {p["key"]: p["lines"] for p in split_jd(sample_job().description_text)}
    assert any("GD&T" in line for line in parts["requirements"])
    assert any("Python" in line for line in parts["nice_to_have"])


# ------------------------------------------------------------------ CV formats
def test_country_lookup():
    assert code_for_country("Germany") == "DE"
    assert code_for_country("United Arab Emirates") == "AE"
    assert code_for_country("Atlantis") == "IN"


def test_us_and_canada_never_show_personal_details_or_photo():
    profile = sample_profile()
    profile.date_of_birth, profile.nationality, profile.visa_status = "14.03.1997", "Indian", "Employment Visa"
    for code in ("US", "CA", "GB", "NL"):
        fmt = get_format(code)
        doc = build(sample_cv(), fmt, STYLES["modern"], profile)
        assert doc.personal_lines == [] and doc.photo is None, code


def test_gulf_and_germany_show_only_details_the_user_gave():
    profile = sample_profile()
    profile.nationality = "Indian"
    de = build(sample_cv(), get_format("DE"), STYLES["classic"], profile)
    assert de.personal_lines == ["Nationality: Indian"]  # no date of birth because none was given
    assert de.signature and "01/2021" in str(de.sections[0][2])  # MM/YYYY dates
    ae = build(sample_cv(), get_format("AE"), STYLES["classic"], profile)
    assert "Nationality: Indian" in ae.personal_lines


def test_page_limit_is_enforced_by_removing_not_adding():
    cv = sample_cv()
    for exp in cv.experience:
        exp.bullets = exp.bullets * 5
    profile = sample_profile()
    _, pdf, doc = render(cv, get_format("US"), STYLES["classic"], profile)
    assert pdf_page_count(pdf) == 1
    assert doc.removed
    rendered_bullets = [b for _, _, payload in doc.sections if isinstance(payload, list)
                        for item in payload if isinstance(item, dict) for b in item.get("bullets", [])]
    original = {b for e in cv.experience for b in e.bullets}
    assert set(rendered_bullets) <= original


def test_every_format_renders():
    for code in FORMATS:
        for style in STYLES:
            docx_bytes, pdf, _ = render(sample_cv(), get_format(code), STYLES[style], sample_profile())
            assert pdf.startswith(b"%PDF") and docx_bytes[:2] == b"PK"


def test_learn_layout_from_sample_docx():
    import docx

    d = docx.Document()
    for line in ["Jane Example", "Date of Birth: 01.01.1990", "PROFILE", "Operations leader.", "WORK EXPERIENCE",
                 "Ops Manager, Example Ltd", "EDUCATION", "MBA", "SKILLS", "Lean, SAP", "LANGUAGES", "English",
                 "REFERENCES", "Available"]:
        d.add_paragraph(line)
    buffer = io.BytesIO()
    d.save(buffer)
    layout = learn_layout("sample.docx", buffer.getvalue(), "Sample")
    assert layout["section_order"][:4] == ["summary", "experience", "education", "skills"]
    assert "dob" in layout["personal_details"]
    fmt, style = resolve_layout({"country_code": "GB", "style": "modern", "custom_layout_id": layout["id"]}, [layout])
    assert fmt.section_order[0] == "summary" and style["headings"]["experience"] == "Work Experience"


def test_formats_do_not_change_facts():
    cv = sample_cv()
    assert check_cv(cv, sample_facts(cv)).passed  # formatting is applied after the guard, on the same data


# ------------------------------------------------------------------ inbox
def _msg(sender, subject, body):
    return InboxMessage("id-" + subject, sender, sender.split("@")[-1].strip(">"), subject,
                        datetime.now(timezone.utc), body)


def test_classify_emails():
    assert classify("Interview invitation – Operations Manager", "Please share your availability")[0] == "interview"
    assert classify("Your application", "Unfortunately we have decided to move forward with other candidates")[0] == "rejection"
    assert classify("Offer of employment", "We are pleased to offer you")[0] == "offer"
    assert classify("Thank you for applying", "We have received your application")[0] == "acknowledgement"
    assert classify("Newsletter", "Our latest products")[0] == "other"


def test_match_and_detect():
    apps = [{"id": 1, "company": "DHL Supply Chain", "company_domain": "", "title": "Operations Manager"},
            {"id": 2, "company": "Maersk", "company_domain": "", "title": "Supply Chain Manager"}]
    msg = _msg("Talent Team <careers@dhl.com>", "Interview: Operations Manager", "Can we schedule a call?")
    matched = match_application(msg, apps)
    assert matched is not None and matched["id"] == 1
    board = _msg("jobs@linkedin.com", "Maersk is hiring", "interview tips")
    assert match_application(board, apps) is None
    found = detect([msg, board], apps)
    assert len(found) == 1 and found[0].suggested_status == "interviewing"


# ------------------------------------------------------------------ result-oriented CV helpers
def test_bullet_coach_leaves_blanks_for_real_numbers():
    from backend.cv.bullet_coach import has_placeholders, suggest

    tip = suggest("Responsible for managing warehouse operations and reducing costs.")
    assert tip and tip.startswith("Managed warehouse operations") and "[__]" in tip
    assert has_placeholders(tip)
    assert suggest("Reduced costs by 14% in 12 months") is None


def test_evidenced_jd_keywords_are_added_to_skills_but_nothing_else():
    from backend.ai.generation import deterministic_tailor

    cv = sample_cv()
    cv.skills.remove("PPAP")  # still written in a Mahindra bullet, so it is evidenced
    job = sample_job(description_text="Requirements\n- PPAP and DFMEA.\n- Kubernetes.\n- 3+ years of experience.")
    tailored = deterministic_tailor(cv, job, None)
    assert "PPAP" in tailored.skills
    assert "Kubernetes" not in tailored.skills
    assert check_cv(tailored, sample_facts(cv)).passed


def test_relocation_line_only_for_countries_without_work_rights():
    from backend.cv.formatted_renderer import relocation_line

    profile = sample_profile()
    profile.relocation_statement = "Open to international relocation with visa sponsorship."
    assert relocation_line(profile, "Germany") == profile.relocation_statement
    assert relocation_line(profile, "India") == ""


def test_strength_check_flags_missing_manager_evidence():
    from backend.cv.strength import assess

    report = assess(sample_cv())
    found = {i["key"]: i["found"] for i in report.items}
    assert found["team_size"] is False and found["pnl"] is False and found["tools"] is True


def test_hybrid_layout_promotions_and_year_dates():
    from backend.cv.schema import ExperienceItem

    cv = sample_cv()
    cv.experience.insert(1, ExperienceItem(employer="Tata AutoComp Systems", title="Junior Design Engineer",
                                           start="Jan 2020", end="Dec 2020",
                                           bullets=["Drew 120 sheet metal parts in SolidWorks.", "Supported trials."]))
    fmt, style = resolve_layout({"country_code": "GB", "style": "modern", "layout": "hybrid",
                                 "date_precision": "year"}, [])
    doc = build(cv, fmt, style, sample_profile())
    keys = [k for k, _, _ in doc.sections]
    assert keys.index("achievements") < keys.index("experience")
    groups: list = next(p for k, _, p in doc.sections if k == "experience")  # type: ignore[assignment]
    assert [r["title"] for r in groups[0]["roles"]] == ["Design Engineer", "Junior Design Engineer"]
    assert groups[0]["dates"] == "2020 – Present"
    original = {b for e in cv.experience for b in e.bullets}
    achievements: list = next(p for k, _, p in doc.sections if k == "achievements")  # type: ignore[assignment]
    assert all(a.rsplit(" (", 1)[0] + "." in original or a.rsplit(" (", 1)[0] in original for a in achievements)


def test_contract_label_only_when_marked_and_guarded():
    cv = sample_cv()
    cv.experience[1].employment_type = "contract"
    doc = build(cv, get_format("GB"), STYLES["modern"], sample_profile())
    groups: list = next(p for k, _, p in doc.sections if k == "experience")  # type: ignore[assignment]
    titles = [r["title"] for g in groups for r in g["roles"]]
    assert titles == ["Design Engineer", "Graduate Engineer Trainee (Contract)"]
    tampered = cv.model_copy(deep=True)
    tampered.experience[0].employment_type = "consulting"
    assert not check_cv(tampered, sample_facts(cv)).passed


def test_mobility_lines_and_progression():
    from backend.cv.formatted_renderer import mobility_lines
    from backend.cv.strength import progression_line

    profile = sample_profile()
    profile.relocation_statement, profile.passport_status = "Ready to relocate internationally.", "Valid Indian Passport"
    assert mobility_lines(profile, "Canada") == ["Ready to relocate internationally.", "Passport: Valid Indian Passport"]
    assert mobility_lines(profile, "India") == []
    cv = sample_cv()
    cv.experience[0].title, cv.experience[1].title = "Operations Manager", "Executive"
    assert "from Executive to Operations Manager" in progression_line(cv, 7.3)
    assert progression_line(sample_cv(), 7.3) == "" or "Design Engineer" in progression_line(sample_cv(), 7.3)


# ------------------------------------------------------------------ any CV -> roles -> priority
def test_roles_suggested_from_any_cv():
    from backend.cv.structurer import structure_cv
    from backend.jobs.role_suggest import suggest_roles
    from tests.conftest import ROOT

    ops = structure_cv((ROOT / "data/samples/sample_operations_cv.txt").read_text(encoding="utf-8"))
    assert ops.experience[1].title == "Assistant Manager - Warehouse" and ops.experience[1].employer == "Snowman Logistics"
    titles = [r["title"] for r in suggest_roles(ops)]
    assert titles[0] == "Operations Manager" and "Warehouse Manager" in titles
    assert not any("Assistant" in t or t == "Executive" for t in titles)
    design = [r["title"] for r in suggest_roles(sample_cv())]
    assert "Mechanical Design Engineer" in design and "Warehouse Manager" not in design


def test_priority_country_list_then_your_roles_then_cv_roles():
    from backend.jobs.matcher import JobPreferences, priority

    prefs = JobPreferences(target_roles=["Operations Manager"], cv_roles=["Warehouse Manager"],
                           preferred_countries=["Germany", "Netherlands"])
    de_yours = priority(sample_job(title="Operations Manager", location="Berlin, Germany", source_job_id="a"), 70, prefs)
    de_cv = priority(sample_job(title="Warehouse Manager", location="Hamburg, Germany", source_job_id="b"), 90, prefs)
    nl_yours = priority(sample_job(title="Operations Manager", location="Amsterdam, Netherlands", source_job_id="c"), 95, prefs)
    us_yours = priority(sample_job(title="Operations Manager", location="Austin, United States", source_job_id="d"), 99, prefs)
    ranks = [de_yours["rank"], de_cv["rank"], nl_yours["rank"], us_yours["rank"]]
    assert ranks == sorted(ranks)
    assert de_yours["role_tier"] == 0 and de_cv["role_tier"] == 1 and not us_yours["in_country_list"]


# ------------------------------------------------------------------ plans
def test_pro_status_and_referral_extension():
    import pytest

    pytest.importorskip("sqlmodel")
    from datetime import datetime, timedelta, timezone

    from backend.models import User
    from backend.services.billing import extend_pro, is_pro

    u = User(email="a@example.com", password_hash="x")
    assert not is_pro(u)
    extend_pro(u, 7, "referral")
    assert is_pro(u) and u.plan_source == "referral"
    first = u.plan_expires_at
    extend_pro(u, 7, "referral")
    assert u.plan_expires_at - first == timedelta(days=7)  # stacks
    u.plan_expires_at = datetime.now(timezone.utc) - timedelta(days=1)
    assert not is_pro(u)
